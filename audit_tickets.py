"""
Ticket audit: compares an MIS Excel file with ticket PDFs and writes an
audited workbook with three sheets: Summary, Reconciliation and
Audited MIS Report.

Each comparable field is shown side by side (MIS | PDF | Match) in the
Reconciliation sheet, with green for Yes/MATCH and red for No/MISMATCH.
"""
import os
import re
import glob
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from audit_rules import (
    names_match_any,
    classes_match,
    bus_classes_match,
    quota_matches,
    stations_match,
    dates_match,
    fare_matches,
    make_status,
    yes_no,
    build_remarks,
    station_remark,
)
from irctc_ers import extract_irctc_ers


# ============================================================
# OUTPUT STRUCTURE
# ============================================================

RECONCILIATION_HEADERS = [
    "Status", "PNR / Ticket No", "Mode", "Created By", "PDF File",
    "Passenger (MIS)", "Passenger (PDF)", "Name Match",
    "Class (MIS)", "Class (PDF)", "Class Match",
    "Quota (MIS)", "Quota (PDF)", "Quota Match",
    "Travel Date (MIS)", "Travel Date (PDF)", "Date Match",
    "From (MIS)", "From (PDF)", "From Match",
    "To (MIS)", "To (PDF)", "To Match",
    "Ticket Amount (MIS)", "Fare/Pax (PDF)", "Fare Match",
    "Remarks",
]

AUDITED_HEADERS = [
    "Status", "Corporate Name", "Invoice No", "Invoice Date",
    "Date Of Booking", "PNR Number", "Mode", "Train_Bus_Flight_Name",
    "Train_Bus_Flight_Number", "Date Of Travel", "Passenger Name",
    "Employee Code", "Travel Request No", "Cost Center", "From", "To",
    "Class", "Quota", "Ticket Amount", "Charges", "GST 18%",
    "Total Amount", "Created By", "VendorName", "Remarks",
]

# Column positions (1-based) in the Reconciliation sheet
RECON_MATCH_COLUMNS = [8, 11, 14, 17, 20, 23, 26]   # Name, Class, Quota, Date, From, To, Fare


# ============================================================
# STYLES
# ============================================================

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")

MATCH_FILL = PatternFill("solid", fgColor="C6EFCE")
MATCH_FONT = Font(bold=True, color="006100")

MISMATCH_FILL = PatternFill("solid", fgColor="FFC7CE")
MISMATCH_FONT = Font(bold=True, color="9C0006")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

GREEN_STATUSES = {"MATCH", "YES"}
RED_STATUSES = {"MISMATCH", "MISSING PDF", "UNRECORDED", "NO"}


# ============================================================
# BASIC HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_text(value):
    value = clean(value).upper()
    value = value.replace("\n", " ").replace("\r", " ").replace("-", " ")
    return re.sub(r"\s+", " ", value).strip()


def normalize_number(value):
    if value is None:
        return None
    text = clean(value).replace(",", "").replace("₹", "")
    text = text.replace("INR", "").replace("Rs.", "").replace("Rs", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    return round(float(match.group()), 2)


DATE_FORMATS = [
    "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%y", "%d/%m/%y",
    "%Y-%m-%d", "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%d-%b-%Y", "%d-%B-%Y",
]


def normalize_date(value):
    """Returns dd-mm-yyyy when the value is a date; otherwise the cleaned text."""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    text = clean(value).split(" ")[0] if clean(value) else ""
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%d-%m-%Y")
        except ValueError:
            pass
    # Full-string attempt for values like "27 Sep 2026" (contains spaces)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(clean(value), fmt).strftime("%d-%m-%Y")
        except ValueError:
            pass
    return clean(value)


def display_date(value):
    """Date as shown in the output: dd-mm-yyyy if it is a date, else the original text."""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    return value if value is not None else ""


def is_railway(mode):
    mode = normalize_text(mode)
    return "RAIL" in mode or "TRAIN" in mode or "IRCTC" in mode


# ============================================================
# EXCEL HELPERS
# ============================================================

def get_header_map(ws):
    result = {}
    for cell in ws[1]:
        if cell.value is not None:
            result[normalize_text(cell.value)] = cell.column
    return result


def get_value(ws, row, header_map, possible_names):
    for name in possible_names:
        column = header_map.get(normalize_text(name))
        if column is not None:
            return ws.cell(row=row, column=column).value
    return None


MIS_FIELDS = {
    "corporate": ["Corporate Name"],
    "invoice_no": ["Invoice No", "Invoice Number"],
    "invoice_date": ["Invoice Date"],
    "booking_date": ["Date Of Booking", "Date of Booking"],
    "pnr": ["PNR Number", "PNR", "Ticket No", "Ticket Number"],
    "mode": ["Mode"],
    "train_name": ["Train_Bus_Flight_Name", "Train/Bus/Flight Name", "Train Name", "Flight Name", "Bus Name"],
    "train_number": ["Train_Bus_Flight_Number", "Train/Bus/Flight Number", "Train Number", "Flight Number", "Bus Number"],
    "travel_date": ["Date Of Travel", "Travel Date"],
    "passenger": ["Passenger Name", "Passenger", "Name"],
    "employee_code": ["Employee Code"],
    "travel_request": ["Travel Request No", "Travel Request Number"],
    "cost_center": ["Cost Center"],
    "from": ["From", "Origin"],
    "to": ["To", "Destination"],
    "class": ["Class", "Travel Class"],
    "quota": ["Quota"],
    "amount": ["Ticket Amount", "Ticket Fare", "Fare", "Amount"],
    "charges": ["Charges"],
    "gst": ["GST 18%", "GST"],
    "total": ["Total Amount"],
    "created_by": ["Created By"],
    "vendor": ["VendorName", "Vendor Name"],
}


# ============================================================
# PDF TEXT
# ============================================================

def extract_pdf_text(pdf_path):
    try:
        reader = PdfReader(pdf_path)
        pages = []
        for page in reader.pages:
            try:
                text = page.extract_text()
                if text:
                    pages.append(text)
            except Exception:
                pass
        return "\n".join(pages)
    except Exception:
        return ""


# ============================================================
# GENERIC PDF EXTRACTION (non-IRCTC tickets)
# ============================================================

NOISE = re.compile(r"SUPPLIER|ADDRESS|STATE NAME|TERMS|INSTRUCTION|CERTIFICATE", re.I)


def clean_field(value):
    """Rejects values that look like label text or boilerplate."""
    value = clean(value)
    if not value or len(value) > 60 or ":" in value or NOISE.search(value):
        return ""
    return value


def extract_after_label(text, labels):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for i, line in enumerate(lines):
        if NOISE.search(line):
            continue
        for label in labels:
            m = re.match(
                r"^" + re.escape(label) + r"(?![A-Za-z/])\s*[:|\-]?\s*(.*)$",
                line,
                re.IGNORECASE,
            )
            if m:
                rest = m.group(1).strip()
                if rest:
                    return rest
                if i + 1 < len(lines):
                    return lines[i + 1]
    return ""


def extract_pnr(text):
    pattern = (
        r"\b(?:PNR|BOOKING ID|TICKET NO)[ \t]*(?:NUMBER|NO|NUM|#|ID)?[ \t]*[:\-]?[ \t]*"
        r"((?=[A-Z0-9]*\d)[A-Z0-9]{5,20})\b"
    )
    m = re.search(pattern, text, re.IGNORECASE)
    return m.group(1) if m else ""


CLASS_CODES = ["1A", "2A", "3A", "SL", "CC", "EC", "2S", "3E", "FC"]
QUOTA_WORDS = ["PREMIUM TATKAL", "TATKAL", "GENERAL", "LADIES", "SENIOR CITIZEN"]


def extract_generic_fields(text):
    upper = normalize_text(text)

    passenger = clean_field(extract_after_label(text, ["Passenger Name", "Passenger"]))

    travel_class = clean_field(extract_after_label(text, ["Class", "Travel Class"]))
    if not travel_class:
        for code in CLASS_CODES:
            if re.search(r"\b" + code + r"\b", upper):
                travel_class = code
                break

    quota = clean_field(extract_after_label(text, ["Quota"]))
    if not quota:
        for word in QUOTA_WORDS:
            if re.search(r"\b" + word + r"\b", upper):
                quota = word.title()
                break

    travel_date = clean_field(extract_after_label(
        text, ["Date Of Travel", "Date of Travel", "Travel Date", "Journey Date", "Date of Journey"]))
    if not travel_date:
        m = re.search(r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b", text)
        travel_date = m.group() if m else ""

    origin = clean_field(extract_after_label(text, ["From", "Origin", "Boarding From", "Source"]))
    destination = clean_field(extract_after_label(text, ["To", "Destination", "Drop", "Arrival"]))

    fare = ""
    for pattern in [
        r"(?:Total Fare|Ticket Fare|Ticket Amount|Amount Payable|Total Amount)\s*[:\-]?\s*(?:₹|INR|Rs\.?)?\s*([\d,]+(?:\.\d+)?)",
    ]:
        found = re.findall(pattern, text, re.IGNORECASE)
        if found:
            fare = found[-1]
            break

    return {
        "pnr": extract_pnr(text),
        "passenger": passenger,
        "passengers": [passenger] if passenger else [],
        "class": travel_class,
        "quota": quota,
        "travel_date": travel_date,
        "from": origin,
        "to": destination,
        "fare": fare,
    }


def extract_pdf_fields(pdf_path):
    text = extract_pdf_text(pdf_path)

    if re.search(r"Electronic Reservation Slip|IRCTC", text, re.IGNORECASE):
        fields = extract_irctc_ers(text)
        fields["mode"] = "Train"
        return fields

    fields = extract_generic_fields(text)
    fields["mode"] = ""
    return fields


# ============================================================
# PDF INDEX
# ============================================================

def build_pdf_index(tickets_dir):
    """
    Returns:
      index     : {PNR: {"file", "path", "fields"}} for the first PDF of each PNR
      pdf_files : every PDF found, sorted
      cache     : {path: fields} so each PDF is parsed once
    """
    pdf_files = sorted(glob.glob(os.path.join(tickets_dir, "**", "*.pdf"), recursive=True))

    index, cache = {}, {}
    for path in pdf_files:
        fields = extract_pdf_fields(path)
        cache[path] = fields

        pnr = normalize_text(fields.get("pnr", ""))
        if pnr and pnr not in index:
            index[pnr] = {
                "file": os.path.basename(path),
                "path": path,
                "fields": fields,
            }

    return index, pdf_files, cache


# ============================================================
# FORMATTING
# ============================================================

def style_status(cell):
    value = normalize_text(cell.value)
    if value in GREEN_STATUSES:
        cell.fill, cell.font = MATCH_FILL, MATCH_FONT
    elif value in RED_STATUSES:
        cell.fill, cell.font = MISMATCH_FILL, MISMATCH_FONT


def format_sheet(ws, status_column=1, match_columns=None):
    # Header
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 35

    # Data
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        style_status(row[status_column - 1])

        for col in match_columns or []:
            cell = row[col - 1]
            value = normalize_text(cell.value)
            if value == "YES":
                cell.fill, cell.font = MATCH_FILL, MATCH_FONT
            elif value == "NO":
                cell.fill, cell.font = MISMATCH_FILL, MISMATCH_FONT

    ws.freeze_panes = "B2"

    # Column widths
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 35)


# ============================================================
# MAIN
# ============================================================

def audit_tickets(excel_path, tickets_dir, output_path,
                  report_date=None, folder_label="tickets"):
    report_date = report_date or datetime.now().strftime("%d-%m-%Y")

    # --------------------------------------------------------
    # LOAD MIS
    # --------------------------------------------------------
    source_ws = load_workbook(excel_path, data_only=False).active
    header_map = get_header_map(source_ws)

    records = []
    pnr_totals = {}

    for row in range(2, source_ws.max_row + 1):
        values = [source_ws.cell(row=row, column=c).value for c in range(1, source_ws.max_column + 1)]
        if not any(v is not None and str(v).strip() for v in values):
            continue

        rec = {name: get_value(source_ws, row, header_map, names) for name, names in MIS_FIELDS.items()}
        rec["pnr"] = normalize_text(rec["pnr"])
        rec["row"] = row
        records.append(rec)

        amount = normalize_number(rec["amount"])
        if rec["pnr"] and amount is not None:
            pnr_totals[rec["pnr"]] = pnr_totals.get(rec["pnr"], 0) + amount

    mis_pnrs = {r["pnr"] for r in records if r["pnr"]}

    # --------------------------------------------------------
    # INDEX PDFs
    # --------------------------------------------------------
    pdf_index, pdf_files, pdf_cache = build_pdf_index(tickets_dir)

    # --------------------------------------------------------
    # RECONCILE EACH MIS ROW
    # --------------------------------------------------------
    recon_rows, audited_rows = [], []
    counts = {"MATCH": 0, "MISMATCH": 0, "MISSING PDF": 0, "UNRECORDED": 0}
    used_pdfs = set()

    for rec in records:
        railway = is_railway(rec["mode"])
        pdf = pdf_index.get(rec["pnr"])

        if pdf is None:
            status = "MISSING PDF"
            counts["MISSING PDF"] += 1
            remarks = f"PDF ticket not found in {folder_label} folder"

            pdf_file = ""
            p_pax, p_class, p_quota = "", "", ""
            p_date, p_from, p_to, p_fare = "", "", "", None

        else:
            used_pdfs.add(pdf["path"])
            pdf_fields = pdf["fields"]
            pdf_file = pdf["file"]

            p_pax = pdf_fields.get("passenger", "")
            p_class = pdf_fields.get("class", "")
            p_quota = pdf_fields.get("quota", "")
            p_date = normalize_date(pdf_fields.get("travel_date", ""))
            p_from = pdf_fields.get("from", "")
            p_to = pdf_fields.get("to", "")
            p_fare = normalize_number(pdf_fields.get("fare", ""))

            pdf_passengers = pdf_fields.get("passengers") or [p_pax]

            name_ok = names_match_any(rec["passenger"], pdf_passengers)
            class_ok = (classes_match if railway else bus_classes_match)(rec["class"], p_class)
            quota_ok = quota_matches(rec["quota"], p_quota, railway)
            date_ok = dates_match(normalize_date(rec["travel_date"]), p_date)
            from_ok = stations_match(rec["from"], p_from)
            to_ok = stations_match(rec["to"], p_to)
            fare_ok = fare_matches(rec["amount"], pdf_fields.get("fare", ""), pnr_totals.get(rec["pnr"]))

            status = make_status([name_ok, class_ok, quota_ok, date_ok, from_ok, to_ok, fare_ok])
            counts[status] += 1

            problems = []
            if not name_ok:
                problems.append(f"Passenger Name Mismatch: Excel '{clean(rec['passenger'])}' vs PDF '{p_pax}'")
            if not class_ok:
                problems.append(f"Class Mismatch: Excel '{clean(rec['class'])}' vs PDF '{p_class}'")
            if not quota_ok:
                problems.append(f"Quota Mismatch: Excel '{clean(rec['quota'])}' vs PDF '{p_quota}'")
            if not date_ok:
                problems.append(
                    f"Travel Date Mismatch: Excel '{normalize_date(rec['travel_date'])}' vs PDF '{p_date}'")
            if not from_ok:
                problems.append(station_remark("Origin", clean(rec["from"]), p_from))
            if not to_ok:
                problems.append(station_remark("Destination", clean(rec["to"]), p_to))
            if not fare_ok:
                problems.append(f"Fare Mismatch: Excel '{clean(rec['amount'])}' vs PDF '{pdf_fields.get('fare', '')}'")
            remarks = build_remarks(problems)

        # Yes/No for each comparison (Missing PDF rows are all "No")
        if pdf is None:
            name_ok = class_ok = quota_ok = date_ok = from_ok = to_ok = fare_ok = False

        name_m, class_m, quota_m = yes_no(name_ok), yes_no(class_ok), yes_no(quota_ok)
        date_m, from_m, to_m, fare_m = yes_no(date_ok), yes_no(from_ok), yes_no(to_ok), yes_no(fare_ok)

        recon_rows.append([
            status, rec["pnr"], clean(rec["mode"]), clean(rec["created_by"]), pdf_file,
            clean(rec["passenger"]), p_pax, name_m,
            clean(rec["class"]), p_class, class_m,
            clean(rec["quota"]), p_quota, quota_m,
            display_date(rec["travel_date"]), p_date, date_m,
            clean(rec["from"]), p_from, from_m,
            clean(rec["to"]), p_to, to_m,
            rec["amount"], p_fare, fare_m,
            remarks,
        ])

        audited_rows.append([
            status, clean(rec["corporate"]), clean(rec["invoice_no"]),
            display_date(rec["invoice_date"]), display_date(rec["booking_date"]),
            rec["pnr"], clean(rec["mode"]), clean(rec["train_name"]), clean(rec["train_number"]),
            display_date(rec["travel_date"]), clean(rec["passenger"]), clean(rec["employee_code"]),
            clean(rec["travel_request"]), clean(rec["cost_center"]),
            clean(rec["from"]), clean(rec["to"]), clean(rec["class"]), clean(rec["quota"]),
            rec["amount"], rec["charges"], rec["gst"], rec["total"],
            clean(rec["created_by"]), clean(rec["vendor"]), remarks,
        ])

    # --------------------------------------------------------
    # UNRECORDED PDFs (in folder, PNR not in MIS)
    # --------------------------------------------------------
    for path in pdf_files:
        if path in used_pdfs:
            continue

        fields = pdf_cache[path]
        pnr = normalize_text(fields.get("pnr", ""))
        if not pnr or pnr in mis_pnrs:
            continue

        counts["UNRECORDED"] += 1
        p_date = normalize_date(fields.get("travel_date", ""))
        p_fare = normalize_number(fields.get("fare", ""))
        remarks = "Found in PDF folder, missing in Excel MIS report"

        recon_rows.append([
            "UNRECORDED", pnr, fields.get("mode", ""), "", os.path.basename(path),
            "", fields.get("passenger", ""), "No",
            "", fields.get("class", ""), "No",
            "", fields.get("quota", ""), "No",
            "", p_date, "No",
            "", fields.get("from", ""), "No",
            "", fields.get("to", ""), "No",
            "", p_fare, "No",
            remarks,
        ])

        audited_rows.append([
            "UNRECORDED", "", "", "", "", pnr, fields.get("mode", ""), "", "",
            p_date, "", "", "", "",
            fields.get("from", ""), fields.get("to", ""), fields.get("class", ""), fields.get("quota", ""),
            p_fare, "", "", "", "", "", remarks,
        ])

    # --------------------------------------------------------
    # WORKBOOK
    # --------------------------------------------------------
    wb = Workbook()
    wb.remove(wb.active)

    summary_ws = wb.create_sheet("Summary")
    recon_ws = wb.create_sheet("Reconciliation")
    audited_ws = wb.create_sheet("Audited MIS Report")

    recon_ws.append(RECONCILIATION_HEADERS)
    for row in recon_rows:
        recon_ws.append(row)

    audited_ws.append(AUDITED_HEADERS)
    for row in audited_rows:
        audited_ws.append(row)

    # Summary
    summary_ws["A1"] = "Reconciliation Summary"
    summary_ws["B1"] = report_date
    summary_ws["A1"].font = Font(bold=True, size=14)
    summary_ws["B1"].font = Font(bold=True)

    summary_ws["A3"] = "Status Metric"
    summary_ws["B3"] = "Count"
    for cell in summary_ws[3]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    metrics = [
        ("Fully Matched (MATCH)", counts["MATCH"], MATCH_FILL),
        ("Discrepancies / Mismatch (MISMATCH)", counts["MISMATCH"], MISMATCH_FILL),
        ("Missing PDF (No PDF Uploaded)", counts["MISSING PDF"], MISMATCH_FILL),
        ("Unrecorded Bookings (In Folder, Missing in MIS)", counts["UNRECORDED"], MISMATCH_FILL),
        ("Total MIS Manifest Rows", len(records), None),
        ("Total PDFs Scanned in Folder", len(pdf_files), None),
    ]
    for i, (label, value, fill) in enumerate(metrics, start=4):
        summary_ws[f"A{i}"] = label
        summary_ws[f"B{i}"] = value
        for col in ("A", "B"):
            summary_ws[f"{col}{i}"].border = THIN_BORDER
            summary_ws[f"{col}{i}"].alignment = Alignment(vertical="center", wrap_text=True)
        if fill:
            summary_ws[f"B{i}"].fill = fill

    summary_ws.column_dimensions["A"].width = 48
    summary_ws.column_dimensions["B"].width = 20

    # Detail sheets
    format_sheet(recon_ws, status_column=1, match_columns=RECON_MATCH_COLUMNS)
    format_sheet(audited_ws, status_column=1)

    wb.save(output_path)
    return output_path


# ============================================================
# COMMAND LINE
# ============================================================

if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True)
    parser.add_argument("--tickets", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--report-date", default=None, help="dd-mm-yyyy shown on Summary")
    args = parser.parse_args()

    out = audit_tickets(
        excel_path=args.excel,
        tickets_dir=args.tickets,
        output_path=args.output,
        report_date=args.report_date,
    )
    print("Audit completed successfully:")
    print(out)
